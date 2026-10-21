"""Rough visual preview of a sheet (layout sanity check, not pixel-exact Excel)."""
import json, sys
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

XLSX, VALUES, SHEET, OUT = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
R1 = int(sys.argv[5]) if len(sys.argv) > 5 else 1
R2 = int(sys.argv[6]) if len(sys.argv) > 6 else 60
C2 = int(sys.argv[7]) if len(sys.argv) > 7 else 14
SCALE = float(sys.argv[8]) if len(sys.argv) > 8 else 1.0

wb = openpyxl.load_workbook(XLSX)
ws = wb[SHEET]
vals = json.load(open(VALUES)).get(SHEET, {})

def fnt(size, bold=False):
    path = "/usr/share/fonts/truetype/dejavu/DejaVuSans%s.ttf" % ("-Bold" if bold else "")
    try:
        return ImageFont.truetype(path, max(6, int(size * 1.28 * SCALE)))
    except Exception:
        return ImageFont.load_default()

def colw(c):
    d = ws.column_dimensions[get_column_letter(c)]
    return (d.width if d.width else 8.43) * 7 + 5

def rowh(r):
    d = ws.row_dimensions[r]
    return (d.height if d.height else 15) * 1.333

xs = {0: 0}
for c in range(1, C2 + 2):
    xs[c] = xs[c - 1] + colw(c)
ys = {R1 - 1: 0}
for r in range(R1, R2 + 2):
    ys[r] = ys[r - 1] + rowh(r)

W = int(xs[C2 + 1] * SCALE) + 4
H = int(ys[R2 + 1] * SCALE) + 4
img = Image.new("RGB", (max(W, 100), max(H, 100)), "white")
d = ImageDraw.Draw(img)

merged = {}
for rng in ws.merged_cells.ranges:
    c1, r1, c2, r2 = range_boundaries(str(rng))
    merged[(r1, c1)] = (c2, r2)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if (r, c) != (r1, c1):
                merged[(r, c)] = None

for r in range(R1, R2 + 1):
    for c in range(1, C2 + 1):
        cell = ws.cell(row=r, column=c)
        x0, x1 = xs[c - 1] * SCALE, xs[c] * SCALE
        y0, y1 = ys[r - 1] * SCALE, ys[r] * SCALE
        anchor = merged.get((r, c), "keep")
        if anchor is None:
            continue
        span = merged.get((r, c))
        if span:
            x1 = xs[min(span[0], C2 + 1)] * SCALE
            y1 = ys[min(span[1], R2 + 1)] * SCALE
        fg = cell.fill.fgColor.rgb if cell.fill and cell.fill.patternType else None
        if fg and isinstance(fg, str) and fg not in ("00000000",):
            d.rectangle([x0, y0, x1, y1], fill="#" + fg[-6:])
        b = cell.border
        has_border = any([getattr(b.top, 'style', None), getattr(b.bottom, 'style', None),
                          getattr(b.left, 'style', None), getattr(b.right, 'style', None)])
        if has_border and r > 1:
            d.rectangle([x0, y0, x1, y1], outline="#DFE4E8")
        v = vals.get(f"{r},{c}")
        if v is None and isinstance(cell.value, str) and not cell.value.startswith("="):
            v = cell.value
        if v is None or v == "":
            continue
        s = str(v).replace("\n", " ")
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            nf = cell.number_format or ""
            if "\u20b1" in nf or "#,##0" in nf:
                s = f"{v:,.2f}" if (".00" in (nf or "")) else f"{v:,.0f}"
            if "%" in (nf or ""):
                s = f"{v*100:.1f}%"
            elif any(k in (nf or "") for k in ("yy", "mmm", "dd")):
                import datetime as _dt
                dtv = _dt.datetime(1899, 12, 30) + _dt.timedelta(days=float(v))
                s = dtv.strftime("%d %b %Y") if "dd" in nf else dtv.strftime("%b %Y")
        size = cell.font.sz or 11
        bold = bool(cell.font.b)
        colour = cell.font.color.rgb if (cell.font.color and isinstance(cell.font.color.rgb, str)) else "000000"
        colour = "#" + (colour[-6:] if len(colour) >= 6 else "000000")
        f = fnt(size, bold)
        tw = d.textlength(s, font=f)
        al = (cell.alignment.horizontal or "left") if cell.alignment else "left"
        if al == "center":
            tx = x0 + max(2, (x1 - x0 - tw) / 2)
        elif al == "right":
            tx = x1 - tw - 3
        else:
            tx = x0 + 4 + (cell.alignment.indent or 0) * 6 if cell.alignment else x0 + 4
        val = (cell.alignment.vertical or "bottom") if cell.alignment else "bottom"
        if val == "top":
            ty = y0 + 2
        elif val == "center":
            ty = y0 + max(1, (y1 - y0 - (size * 1.28 * SCALE)) / 2)
        else:
            ty = y1 - (size * 1.28 * SCALE) - 2
        d.text((tx, ty), s, font=f, fill=colour)
img.save(OUT)
print("wrote", OUT, img.size)
