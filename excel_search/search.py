from datetime import date, datetime
from pathlib import Path
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import xlrd

SUPPORTED = {'.xlsx', '.xlsm', '.xls'}


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    return str(value)


def normalize(value):
    return ' '.join(unicodedata.normalize('NFKC', text(value)).casefold().split())


def sheets(path):
    """Yield streaming rows; never execute Excel macros or evaluate formulas."""
    if path.suffix.lower() == '.xls':
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                yield sheet.name, (sheet.row_values(i) for i in range(sheet.nrows))
        finally:
            book.release_resources()
    else:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            book.close()


def search_files(files, queries, exact=False, header_row=1, limit=1000):
    needles = [(q, normalize(q)) for q in queries if normalize(q)]
    results, errors = [], []
    scanned = 0
    for path, display_name in files:
        try:
            for sheet_name, rows in sheets(path):
                headers = []
                for number, row in enumerate(rows, 1):
                    if number == header_row:
                        headers = [text(v) for v in row]
                    if number <= header_row:
                        continue
                    values = [text(v) for v in row]
                    matched_columns, matched_queries = [], set()
                    for col, value in enumerate(values):
                        normalized = normalize(value)
                        hits = [q for q, needle in needles if (normalized == needle if exact else needle in normalized)]
                        if hits:
                            matched_columns.append(col)
                            matched_queries.update(hits)
                    if matched_columns:
                        if len(results) >= limit:
                            return dict(results=results, errors=errors, scanned_files=scanned, truncated=True)
                        results.append(dict(file=display_name, sheet=sheet_name, row=number,
                            matched_queries=sorted(matched_queries), matched_columns=matched_columns,
                            cells=[dict(column=get_column_letter(i + 1), label=(headers[i] if i < len(headers) and headers[i] else f'Column {get_column_letter(i + 1)}'), value=v)
                                   for i, v in enumerate(values) if v or i in matched_columns]))
            scanned += 1
        except Exception as exc:
            errors.append(dict(file=display_name, message=f'Could not read workbook ({type(exc).__name__}). It may be damaged, encrypted, or unsupported.'))
    return dict(results=results, errors=errors, scanned_files=scanned, truncated=False)
